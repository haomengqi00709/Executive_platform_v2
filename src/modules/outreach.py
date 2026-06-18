"""
Outreach Tool — batch-generate personalized email drafts from a OneDrive folder.

Scenario: user attends a conference, collects contacts (business cards, attendee lists,
CSV exports). They drop the files into a OneDrive folder. This tool reads each file,
extracts contact info, and creates a personalized Outlook Draft for every contact.

NOT a section — it's an on-demand personalized tool. Output goes to data_dir/outreach/
(not data_dir/results/, which is sections-only).

Entry point: run(graph, ai, data_dir, settings, context_note, folder)
"""
import csv
import io
import json
import os
import re
import tempfile
import time
from datetime import datetime, timezone
from pathlib import Path

try:
    import openpyxl
except ImportError:
    openpyxl = None

from src.graph import GraphClient, _ensure_html
from src.ai import AIClient
from src.modules.profile import load_profile_context, get_user_signature, append_signature_to_body


_IMAGE_EXTS = {".jpg", ".jpeg", ".png", ".gif", ".webp"}
_PDF_EXTS   = {".pdf"}
_CSV_EXTS   = {".csv"}
_XLSX_EXTS  = {".xlsx", ".xls"}

# Hard ceiling on real sends per bulk run. Drafts are unbounded (harmless until
# the user hits send in Outlook); real sends go out of the user's own mailbox,
# where M365 throttles and the domain's reputation is on the line — so a single
# "Send now" batch is capped. The /api/outreach/bulk route re-checks this.
BULK_SEND_CAP = 50


def _load_state(data_dir: Path) -> dict:
    f = Path(data_dir) / "outreach.json"
    if f.exists():
        try:
            return json.loads(f.read_text())
        except Exception:
            pass
    return {"processed_files": {}, "history": []}


def _save_state(data_dir: Path, state: dict) -> None:
    f = Path(data_dir) / "outreach.json"
    tmp = f.with_suffix(".tmp")
    tmp.write_text(json.dumps(state, indent=2, ensure_ascii=False))
    tmp.replace(f)


def _save_result(data_dir: Path, result: dict) -> None:
    out_dir = Path(data_dir) / "outreach"
    out_dir.mkdir(parents=True, exist_ok=True)
    f = out_dir / "results.json"
    tmp = f.with_suffix(".tmp")
    tmp.write_text(json.dumps(result, indent=2, ensure_ascii=False))
    tmp.replace(f)


# ── Contact extraction by file type ─────────────────────────────────────────

def _extract_from_image(ai: AIClient, file_bytes: bytes, ext: str) -> list[dict]:
    """Gemini Vision: extract single contact from a business card photo."""
    mime_map = {".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".png": "image/png",
                ".gif": "image/gif", ".webp": "image/webp"}
    mime = mime_map.get(ext, "image/jpeg")

    with tempfile.NamedTemporaryFile(suffix=ext, delete=False) as f:
        f.write(file_bytes)
        tmp_path = f.name

    try:
        from google.genai import types as genai_types
        uploaded = ai.client.files.upload(file=tmp_path)
        while uploaded.state.name == "PROCESSING":
            time.sleep(2)
            uploaded = ai.client.files.get(name=uploaded.name)

        prompt = """This is a business card photo. Extract the contact info.
Return JSON: {"name": "...", "email": "...", "company": "...", "role": "..."}
If a field is missing, use empty string. If this is NOT a business card, return {"name": ""}."""

        response = ai.client.models.generate_content(
            model=ai.model,
            contents=[
                genai_types.Part.from_uri(file_uri=uploaded.uri, mime_type=mime),
                prompt,
            ],
        )
        ai.record_external_usage(response)
        raw = (response.text or "").strip()
        if raw.startswith("```"):
            raw = raw.split("\n", 1)[-1].rsplit("```", 1)[0]
        match = re.search(r"\{.*\}", raw, re.DOTALL)
        if match:
            data = json.loads(match.group())
            return [data] if data.get("email") else []
    except Exception as e:
        print(f"[Outreach] Image extract failed: {e}")
    finally:
        try:
            os.unlink(tmp_path)
        except Exception:
            pass
    return []


def _extract_from_pdf(ai: AIClient, file_bytes: bytes) -> list[dict]:
    """Gemini multimodal: extract list of contacts from a PDF (attendee list, roster, etc.)."""
    with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as f:
        f.write(file_bytes)
        tmp_path = f.name

    try:
        from google.genai import types as genai_types
        uploaded = ai.client.files.upload(file=tmp_path)
        while uploaded.state.name == "PROCESSING":
            time.sleep(2)
            uploaded = ai.client.files.get(name=uploaded.name)

        prompt = """This PDF likely contains a list of contacts (conference attendees, roster, exported list).
Extract every contact you can find.
Return JSON: {"contacts": [{"name": "...", "email": "...", "company": "...", "role": "..."}, ...]}
Skip rows without an email. If no contacts found, return {"contacts": []}."""

        response = ai.client.models.generate_content(
            model=ai.model,
            contents=[
                genai_types.Part.from_uri(file_uri=uploaded.uri, mime_type="application/pdf"),
                prompt,
            ],
        )
        ai.record_external_usage(response)
        raw = (response.text or "").strip()
        if raw.startswith("```"):
            raw = raw.split("\n", 1)[-1].rsplit("```", 1)[0]
        match = re.search(r"\{.*\}", raw, re.DOTALL)
        if match:
            data = json.loads(match.group())
            return [c for c in data.get("contacts", []) if c.get("email")]
    except Exception as e:
        print(f"[Outreach] PDF extract failed: {e}")
    finally:
        try:
            os.unlink(tmp_path)
        except Exception:
            pass
    return []


def _extract_from_csv(file_bytes: bytes) -> list[dict]:
    """Parse CSV. Try to map columns to name/email/company/role by header name."""
    contacts = []
    try:
        text = file_bytes.decode("utf-8-sig", errors="replace")
        reader = csv.DictReader(io.StringIO(text))
        _known_keys = {"email", "e-mail", "email address", "name", "first name", "last name",
                       "company", "organization", "org", "role", "title", "job title"}
        for row in reader:
            normalized = {k.lower().strip(): (v or "").strip() for k, v in row.items() if k}
            email = (normalized.get("email") or normalized.get("e-mail")
                     or normalized.get("email address") or "")
            if not email or "@" not in email:
                continue
            # Any column NOT in known_keys becomes free-form context (notes, comments, etc.)
            extras = [f"{k}: {v}" for k, v in normalized.items() if k not in _known_keys and v]
            contacts.append({
                "name":    normalized.get("name") or f"{normalized.get('first name','')} {normalized.get('last name','')}".strip(),
                "email":   email,
                "company": normalized.get("company") or normalized.get("organization") or normalized.get("org") or "",
                "role":    normalized.get("role") or normalized.get("title") or normalized.get("job title") or "",
                "notes":   " | ".join(extras),
            })
    except Exception as e:
        print(f"[Outreach] CSV parse failed: {e}")
    return contacts


def _extract_from_xlsx(file_bytes: bytes) -> list[dict]:
    """Parse Excel via openpyxl."""
    if openpyxl is None:
        return []
    contacts = []
    try:
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            f.write(file_bytes)
            tmp_path = f.name
        wb = openpyxl.load_workbook(tmp_path, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(h or "").lower().strip() for h in rows[0]]
        _known_keys = {"email", "e-mail", "email address", "name", "first name", "last name",
                       "company", "organization", "org", "role", "title", "job title"}
        for r in rows[1:]:
            row = dict(zip(headers, [str(v or "").strip() for v in r]))
            email = (row.get("email") or row.get("e-mail") or row.get("email address") or "")
            if not email or "@" not in email:
                continue
            extras = [f"{k}: {v}" for k, v in row.items() if k not in _known_keys and v]
            contacts.append({
                "name":    row.get("name") or f"{row.get('first name','')} {row.get('last name','')}".strip(),
                "email":   email,
                "company": row.get("company") or row.get("organization") or "",
                "role":    row.get("role") or row.get("title") or "",
                "notes":   " | ".join(extras),
            })
        try:
            os.unlink(tmp_path)
        except Exception:
            pass
    except Exception as e:
        print(f"[Outreach] XLSX parse failed: {e}")
    return contacts


# ── Draft generation ────────────────────────────────────────────────────────

def _generate_draft(ai: AIClient, contact: dict, context_note: str,
                    display_name: str, business_context: str,
                    writing_style: str, contact_style: str = "",
                    contact_history: str = "", intent: str = "") -> dict | None:
    """Generate subject + body for one contact. Returns None on failure.

    business_context already includes the user's Personal Profile, Business Profile,
    and Market Segments (merged by load_profile_context).

    Body comes back WITHOUT a signature — caller appends via
    profile.append_signature_to_body(body, get_user_signature(settings)).
    Centralising the append means every draft (this outreach worker, the
    bot's create_reply_draft, the dashboard's draft_generate) gets the
    same signature treatment without each path duplicating logic.
    """
    bc_block = f"User profile (personal + business + market):\n{business_context.strip()}\n\n" if business_context.strip() else ""
    # Per-contact style (extracted from the user's replies to THIS person) beats
    # the user's generic writing_style when available — that's what makes a bulk
    # send read as personal rather than templated.
    effective_style = (contact_style or "").strip() or (writing_style or "").strip()
    style_block = f"Writing style to match:\n{effective_style}\n\n" if effective_style else ""
    history_block = f"Prior relationship/history with this contact (reference it specifically, don't invent):\n{contact_history.strip()}\n\n" if (contact_history or "").strip() else ""
    ctx_block = f"How the user knows this contact:\n{context_note.strip()}\n\n" if context_note.strip() else ""
    notes_block = f"Specific notes from past interaction with this contact:\n{contact.get('notes', '')}\n\n" if (contact.get('notes') or '').strip() else ""
    # The point of the email — in bulk mode this is what the user typed once and
    # wants every recipient's version to be about, rewritten in their voice.
    intent_block = f"Core message this email MUST convey (this is the point — build the email around it, in the user's voice):\n{intent.strip()}\n\n" if (intent or "").strip() else ""

    prompt = f"""You are drafting a personal outreach email on behalf of {display_name}.

{bc_block}{style_block}{history_block}{ctx_block}Contact info:
  Name: {contact.get('name', '')}
  Email: {contact.get('email', '')}
  Company: {contact.get('company', '')}
  Role: {contact.get('role', '')}

{notes_block}{intent_block}Write a personalized outreach email with these requirements:

1. STRUCTURE — 2-3 short paragraphs:
   • Opening: greeting on its own line, then 1-2 sentences referencing the meeting context (or how you know them).
   • Middle: 1-2 sentences connecting their company/role/notes to what the user offers. If "notes from past interaction" exist, weave them in specifically — that's the most personal hook.
   • Closing: 1 sentence with a clear next step (a call, a quick chat, a question).

2. FORMAT — return the body as HTML. Each paragraph wrapped in <p>...</p>. Do NOT use <div>, headers, lists, or styles.

3. TONE — warm but specific. Avoid generic phrases like "great to connect" without a follow-up specific. No marketing jargon.

4. LENGTH — keep total body 60-100 words.

5. DO NOT write a sign-off or signature — the user's real signature (extracted from their actual sent emails) is appended automatically. Anything you write at the end will look like a duplicate.

Return JSON only: {{"subject": "...", "body": "<p>...</p><p>...</p><p>...</p>"}}"""

    try:
        raw = ai.generate(prompt).strip()
        if raw.startswith("```"):
            raw = raw.split("\n", 1)[-1].rsplit("```", 1)[0]
        match = re.search(r"\{.*\}", raw, re.DOTALL)
        if match:
            data = json.loads(match.group())
            if data.get("subject") and data.get("body"):
                return {"subject": data["subject"], "body": data["body"]}
    except Exception as e:
        print(f"[Outreach] Draft gen failed for {contact.get('email')}: {e}")
    return None


def _fill_template(subject: str, body: str, contact: dict) -> tuple:
    """Substitute {name}/{first_name}/{company}/{role} placeholders for a
    'same message to everyone' (non-personalized) bulk send. Missing name
    falls back to 'there' so we never emit 'Hi ,'."""
    name = (contact.get("name") or "").strip()
    first = name.split()[0] if name else ""
    repl = {
        "{name}":       name or "there",
        "{first_name}": first or "there",
        "{company}":    (contact.get("company") or "").strip(),
        "{role}":       (contact.get("role") or "").strip(),
    }

    def _sub(s: str) -> str:
        for k, v in repl.items():
            s = (s or "").replace(k, v)
        return s

    return _sub(subject), _sub(body)


# ── CRM bulk email: two-phase (generate preview → user edits → commit) ──────
#
# Phase 1 generate_bulk() produces editable previews and touches NOTHING in
# Outlook. Phase 2 commit_bulk() takes the user-edited previews and only THEN
# saves drafts / sends. This is the human-review gate the one-shot path lacked.

_BULK_SKILL_FILE = Path(__file__).parent.parent / "skills" / "bulk_email" / "skill.md"


def _generate_bulk_email(ai: AIClient, contact: dict, intent: str,
                         display_name: str, business_context: str,
                         user_instruction: str = "") -> dict | None:
    """Intent-led counterpart to _generate_draft for CRM bulk email: build the
    email AROUND the user's message (intent), adapt tone/structure to the
    contact's relationship status, use the contact's own writing_style + summary.
    Does NOT impose the cold-outreach 'how we met / what I offer / book a call'
    skeleton. Body comes back WITHOUT a signature (commit_bulk appends it).
    Returns None on failure — never silently falls back to cold outreach."""
    skill_text = _BULK_SKILL_FILE.read_text() if _BULK_SKILL_FILE.exists() else ""
    if not skill_text:
        print("[Outreach] bulk_email/skill.md missing — cannot generate intent-led email")
        return None

    prompt = (
        skill_text
        .replace("{display_name}",     display_name or "the executive")
        .replace("{business_context}", (business_context or "").strip() or "(no profile available)")
        .replace("{intent}",           (intent or "").strip())
        .replace("{contact_name}",     contact.get("name", "") or "")
        .replace("{contact_company}",  contact.get("company", "") or "")
        .replace("{contact_role}",     contact.get("role", "") or "")
        .replace("{contact_status}",   contact.get("status", "") or "other")
        .replace("{contact_style}",    (contact.get("writing_style", "") or "").strip() or "(no per-contact style on file)")
        .replace("{contact_history}",  (contact.get("summary", "") or "").strip() or "(no relationship notes on file)")
        .replace("{user_instruction}", (user_instruction or "").strip() or "(none)")
    )
    try:
        raw = ai.generate(prompt).strip()
        if raw.startswith("```"):
            raw = raw.split("\n", 1)[-1].rsplit("```", 1)[0]
        match = re.search(r"\{.*\}", raw, re.DOTALL)
        if match:
            data = json.loads(match.group())
            if data.get("subject") and data.get("body"):
                return {"subject": data["subject"], "body": data["body"]}
    except Exception as e:
        print(f"[Outreach] Bulk email gen failed for {contact.get('email')}: {e}")
    return None


def _save_preview(data_dir: Path, preview: dict) -> None:
    out_dir = Path(data_dir) / "outreach"
    out_dir.mkdir(parents=True, exist_ok=True)
    f = out_dir / "preview.json"
    tmp = f.with_suffix(".tmp")
    tmp.write_text(json.dumps(preview, indent=2, ensure_ascii=False))
    tmp.replace(f)


def generate_bulk(ai: AIClient, data_dir: Path, settings: dict,
                  emails: list, subject: str = "", body: str = "",
                  personalize: bool = True, context_note: str = "",
                  write_files: bool = True, progress=None) -> dict:
    """Phase 1 — generate per-contact email previews. No Outlook calls at all
    (no draft, no send). Writes outreach/preview.json incrementally so the UI
    can poll, then render each item editable before the user commits.

    Bodies come back WITHOUT a signature — commit_bulk appends it, so the user
    edits clean text and the signature stays consistent across the batch."""
    def log(msg: str):
        print(f"[Outreach] {msg}")
        if progress:
            progress(msg)

    from src.modules.crm import find_contacts_by_emails
    data_dir = Path(data_dir)
    display_name     = settings.get("display_name", "the executive")
    business_context = load_profile_context(data_dir) if data_dir else ""
    use_template     = not personalize
    instruction_path = data_dir / "instructions" / "bulk_email.md"
    user_instruction = instruction_path.read_text().strip() if instruction_path.exists() else ""

    contacts = find_contacts_by_emails(data_dir, emails)
    items: list = []
    skipped: list = []
    log(f"Generating {len(contacts)} preview(s) (personalize={personalize})")

    def _preview(done: bool) -> dict:
        return {
            "status": "fresh" if done else "running",
            "last_run": datetime.now(timezone.utc).isoformat(),
            "personalize": bool(personalize),
            "total": len(contacts),
            "context_note": context_note,
            "items": items,
            "skipped": skipped,
        }

    if write_files:
        _save_preview(data_dir, _preview(done=False))

    for i, contact in enumerate(contacts):
        email = (contact.get("email") or "").strip()
        if not email or "@" not in email:
            skipped.append({"reason": "no_email", "contact": contact})
            continue

        if use_template:
            subj, bod = _fill_template(subject, body, contact)
            if not subj.strip() or not bod.strip():
                skipped.append({"reason": "empty_template", "contact": contact})
                continue
        else:
            draft = _generate_bulk_email(
                ai, contact, intent=body,
                display_name=display_name, business_context=business_context,
                user_instruction=user_instruction,
            )
            if not draft:
                skipped.append({"reason": "draft_gen_failed", "contact": contact})
                continue
            # User-typed subject wins (consistent line, {name}/{company} filled);
            # otherwise the AI's per-contact subject.
            subj = _fill_template(subject, "", contact)[0] if subject.strip() else draft["subject"]
            bod = draft["body"]

        items.append({
            "to": email,
            "name": contact.get("name", ""),
            "company": contact.get("company", ""),
            "subject": subj,
            "body": bod,
        })
        if write_files and (i + 1) % 3 == 0:
            _save_preview(data_dir, _preview(done=False))

    preview = _preview(done=True)
    if write_files:
        _save_preview(data_dir, preview)
    log(f"Preview ready — {len(items)} generated, {len(skipped)} skipped")
    return preview


def commit_bulk(graph: GraphClient, data_dir: Path, settings: dict,
                items: list, send: bool = False, write_files: bool = True,
                progress=None) -> dict:
    """Phase 2 — commit the (user-edited) previews. Appends the signature, then
    saves to Outlook Drafts (send=False, the default) or sends for real
    (send=True, capped). `items` = [{to, name, company, subject, body}].
    Writes outreach/results.json (polled by the UI)."""
    def log(msg: str):
        print(f"[Outreach] {msg}")
        if progress:
            progress(msg)

    data_dir = Path(data_dir)
    user_signature = get_user_signature(settings)
    items = items or []

    if send and len(items) > BULK_SEND_CAP:
        result = {
            "status": "not_run", "last_run": datetime.now(timezone.utc).isoformat(),
            "mode": "bulk", "send": True,
            "error": (f"Send batch too large: {len(items)} recipients > cap "
                      f"{BULK_SEND_CAP}. Reduce the selection or save as drafts instead."),
            "drafts_created": [], "files_processed": [], "contacts_skipped": [], "errors": [],
            "summary": {"drafts": 0, "sent": 0, "files": 0, "skipped": 0, "errors": 0},
        }
        if write_files:
            _save_result(data_dir, result)
        return result

    drafts_created: list = []
    errors: list = []

    def _result(done: bool) -> dict:
        sent_n = sum(1 for d in drafts_created if d.get("sent"))
        return {
            "status": "fresh" if done else "running",
            "last_run": datetime.now(timezone.utc).isoformat(),
            "mode": "bulk", "send": bool(send), "total": len(items),
            "drafts_created": drafts_created, "files_processed": [],
            "contacts_skipped": [], "errors": errors,
            "summary": {
                "drafts": len(drafts_created) - sent_n, "sent": sent_n,
                "files": 0, "skipped": 0, "errors": len(errors),
            },
        }

    if write_files:
        _save_result(data_dir, _result(done=False))

    for i, it in enumerate(items):
        email = (it.get("to") or "").strip()
        subj = (it.get("subject") or "")
        bod = (it.get("body") or "")
        if not email or "@" not in email or not subj.strip() or not bod.strip():
            errors.append({"contact": email, "error": "missing to / subject / body"})
            continue
        final_body = append_signature_to_body(bod, user_signature)
        try:
            if send:
                graph.send_mail(to=email, subject=subj, html=_ensure_html(final_body))
                web_link = ""
            else:
                gr = graph.create_draft(subject=subj, body=final_body, to=email)
                web_link = gr.get("webLink", "")
            drafts_created.append({
                "to": email, "name": it.get("name", ""), "company": it.get("company", ""),
                "subject": subj, "web_link": web_link, "sent": bool(send),
                "source_file": "CRM:bulk",
            })
            log(f"  ✅ {'Sent' if send else 'Draft'}: {email}")
        except Exception as e:
            errors.append({"contact": email, "error": str(e)})

        if write_files and (i + 1) % 5 == 0:
            _save_result(data_dir, _result(done=False))
        if send:
            time.sleep(0.5)  # gentle pacing — let M365 throttling breathe

    result = _result(done=True)
    if write_files:
        _save_result(data_dir, result)
    log(f"Commit done — {len(drafts_created)} created "
        f"({sum(1 for d in drafts_created if d.get('sent'))} sent), {len(errors)} errors")
    return result


# ── Main entry ─────────────────────────────────────────────────────────────

def run(graph: GraphClient, ai: AIClient, data_dir: Path, settings: dict,
        context_note: str = "", folder: str = "",
        tag: str = "", recent_hours: int = 0, progress=None) -> dict:
    """
    Generate Outlook drafts for a batch of contacts. Three modes:

      folder       — scan OneDrive folder for files (cards/csv/xlsx/pdf), extract contacts, draft.
      tag          — pull contacts from CRM that have this tag, draft.
      recent_hours — pull contacts from CRM added in the last N hours, draft.

    Defaults to folder mode if none specified. context_note is injected into each draft prompt.
    (CRM multi-select bulk email is a separate two-phase flow — see generate_bulk / commit_bulk.)
    """
    def log(msg: str):
        print(f"[Outreach] {msg}")
        if progress:
            progress(msg)

    data_dir = Path(data_dir)
    display_name     = settings.get("display_name", "the executive")
    business_context = load_profile_context(data_dir) if data_dir else ""
    writing_style    = settings.get("writing_style_note", "")
    # Resolved once for the whole run — same signature appended to every
    # generated draft. Empty string for users who haven't completed
    # signature extraction yet; append helper passes the body through
    # unchanged in that case.
    user_signature   = get_user_signature(settings)

    drafts_created = []
    contacts_skipped = []
    files_processed = []
    errors = []

    # ── Mode 1: CRM by tag or recency ─────────────────────────────────
    if tag or recent_hours:
        from src.modules.crm import find_contacts_by_tag, find_contacts_added_since
        crm_contacts = (find_contacts_by_tag(data_dir, tag) if tag
                        else find_contacts_added_since(data_dir, recent_hours))
        log(f"CRM lookup ({'tag='+tag if tag else 'recent_hours='+str(recent_hours)}) → {len(crm_contacts)} contacts")

        for contact in crm_contacts:
            email = (contact.get("email") or "").strip()
            if not email or "@" not in email:
                contacts_skipped.append({"reason": "no_email", "contact": contact})
                continue
            draft = _generate_draft(ai, contact, context_note,
                                    display_name, business_context, writing_style)
            if not draft:
                contacts_skipped.append({"reason": "draft_gen_failed", "contact": contact})
                continue
            final_body = append_signature_to_body(draft["body"], user_signature)
            try:
                gr = graph.create_draft(subject=draft["subject"], body=final_body, to=email)
                drafts_created.append({
                    "to": email, "name": contact.get("name", ""),
                    "company": contact.get("company", ""),
                    "subject": draft["subject"],
                    "web_link": gr.get("webLink", ""),
                    "source_file": f"CRM:{tag or 'recent'}",
                })
                log(f"  ✅ Draft created: {email}")
            except Exception as e:
                errors.append({"contact": email, "error": str(e)})

        result = {
            "status": "fresh",
            "last_run": datetime.now(timezone.utc).isoformat(),
            "mode": "crm",
            "source": tag or f"recent_{recent_hours}h",
            "context_note": context_note,
            "drafts_created": drafts_created,
            "files_processed": [],
            "contacts_skipped": contacts_skipped,
            "errors": errors,
            "summary": {
                "drafts": len(drafts_created),
                "files": 0,
                "skipped": len(contacts_skipped),
                "errors": len(errors),
            },
        }
        _save_result(data_dir, result)
        log(f"Done — {len(drafts_created)} drafts, {len(contacts_skipped)} skipped")
        return result

    # ── Mode 2: OneDrive folder (original) ─────────────────────────────
    folder = (folder or settings.get("outreach_folder", "")).strip().strip("/")
    if not folder:
        return {
            "status": "not_run",
            "error": "No folder/tag/recent_hours specified.",
            "drafts_created": 0,
            "files_processed": 0,
            "contacts_skipped": 0,
        }

    log(f"Scanning OneDrive folder: {folder}")
    try:
        items = graph.list_drive_folder(folder)
    except Exception as e:
        return {
            "status": "not_run",
            "error": f"Folder not found or inaccessible: {e}",
            "drafts_created": 0,
            "files_processed": 0,
            "contacts_skipped": 0,
        }

    state = _load_state(data_dir)
    processed = state.get("processed_files", {})

    for item in items:
        if item.get("folder"):
            continue  # skip subfolders
        name = item.get("name", "")
        item_id = item.get("id", "")
        ext = Path(name).suffix.lower()

        if item_id in processed:
            continue  # already handled

        log(f"  Processing: {name}")

        try:
            file_bytes = graph.download_drive_item(item_id)
        except Exception as e:
            errors.append({"file": name, "error": f"download failed: {e}"})
            continue

        # Extract contacts based on file type
        if ext in _IMAGE_EXTS:
            contacts = _extract_from_image(ai, file_bytes, ext)
        elif ext in _PDF_EXTS:
            contacts = _extract_from_pdf(ai, file_bytes)
        elif ext in _CSV_EXTS:
            contacts = _extract_from_csv(file_bytes)
        elif ext in _XLSX_EXTS:
            contacts = _extract_from_xlsx(file_bytes)
        else:
            continue  # unknown type, skip silently

        files_processed.append({"file": name, "type": ext, "contacts_found": len(contacts)})

        for contact in contacts:
            email = (contact.get("email") or "").strip()
            if not email or "@" not in email:
                contacts_skipped.append({"reason": "no_email", "file": name, "contact": contact})
                continue

            draft = _generate_draft(
                ai, contact, context_note,
                display_name, business_context, writing_style,
            )
            if not draft:
                contacts_skipped.append({"reason": "draft_gen_failed", "file": name, "contact": contact})
                continue

            final_body = append_signature_to_body(draft["body"], user_signature)
            try:
                result = graph.create_draft(
                    subject=draft["subject"],
                    body=final_body,
                    to=email,
                )
                drafts_created.append({
                    "to": email,
                    "name": contact.get("name", ""),
                    "company": contact.get("company", ""),
                    "subject": draft["subject"],
                    "web_link": result.get("webLink", ""),
                    "source_file": name,
                })
                log(f"    ✅ Draft created: {email}")
            except Exception as e:
                errors.append({"contact": email, "error": f"create_draft failed: {e}"})

        processed[item_id] = {
            "filename": name,
            "processed_at": datetime.now(timezone.utc).isoformat(),
        }

    state["processed_files"] = processed
    state.setdefault("history", []).append({
        "run_at": datetime.now(timezone.utc).isoformat(),
        "folder": folder,
        "context_note": context_note,
        "drafts_created": len(drafts_created),
        "files_processed": len(files_processed),
    })
    state["history"] = state["history"][-50:]  # keep last 50 runs
    _save_state(data_dir, state)

    result = {
        "status": "fresh",
        "last_run": datetime.now(timezone.utc).isoformat(),
        "folder": folder,
        "context_note": context_note,
        "drafts_created": drafts_created,
        "files_processed": files_processed,
        "contacts_skipped": contacts_skipped,
        "errors": errors,
        "summary": {
            "drafts": len(drafts_created),
            "files": len(files_processed),
            "skipped": len(contacts_skipped),
            "errors": len(errors),
        },
    }
    _save_result(data_dir, result)
    log(f"Done — {len(drafts_created)} drafts, {len(files_processed)} files, {len(contacts_skipped)} skipped")
    return result
